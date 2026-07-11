from pathlib import Path
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "redesign_proposal_assets"
DOCX_PATH = ROOT / "docs" / "本地历史导览小程序改版方案与界面图.docx"

W, H = 900, 1600
PHONE_BG = "#F6F4EF"
INK = "#18212F"
MUTED = "#667085"
GREEN = "#15937F"
GREEN_DARK = "#1F6F68"
GREEN_LIGHT = "#EDF4F2"
BORDER = "#E4E0D8"
WHITE = "#FFFFFF"
GOLD = "#B7791F"
BLUE = "#2E74B5"
RED = "#B3261E"


def font(size, weight="regular"):
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc" if weight == "bold" else "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/ArialHB.ttc",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


F = {
    "xs": font(22),
    "sm": font(26),
    "body": font(30),
    "body_b": font(30, "bold"),
    "h3": font(34, "bold"),
    "h2": font(42, "bold"),
    "h1": font(52, "bold"),
    "tab": font(21),
}


def rounded(draw, box, radius=28, fill=WHITE, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text(draw, xy, s, fill=INK, f=None, anchor=None):
    draw.text(xy, s, fill=fill, font=f or F["body"], anchor=anchor)


def fit_text(draw, xy, s, max_width, fill=INK, f=None, line_gap=8, max_lines=None):
    f = f or F["body"]
    line = ""
    lines = []
    for ch in s:
        test = line + ch
        if draw.textlength(test, font=f) <= max_width or not line:
            line = test
        else:
            lines.append(line)
            line = ch
    if line:
        lines.append(line)
    if max_lines and len(lines) > max_lines:
        lines = lines[:max_lines]
        while draw.textlength(lines[-1] + "...", font=f) > max_width and lines[-1]:
            lines[-1] = lines[-1][:-1]
        lines[-1] += "..."
    x, y = xy
    for line in lines:
        draw.text((x, y), line, fill=fill, font=f)
        y += f.size + line_gap
    return y


def icon_camera(draw, cx, cy, color=GREEN):
    rounded(draw, (cx - 34, cy - 24, cx + 34, cy + 30), 14, fill=None, outline=color, width=5)
    draw.rectangle((cx - 18, cy - 36, cx + 18, cy - 24), fill=color)
    draw.ellipse((cx - 17, cy - 9, cx + 17, cy + 25), outline=color, width=5)


def icon_person(draw, cx, cy, color=GREEN):
    draw.ellipse((cx - 18, cy - 36, cx + 18, cy), outline=color, width=5)
    draw.arc((cx - 42, cy - 2, cx + 42, cy + 72), 200, -20, fill=color, width=5)


def icon_route(draw, cx, cy, color=GREEN):
    pts = [(cx - 40, cy + 28), (cx - 10, cy - 22), (cx + 22, cy + 18), (cx + 48, cy - 32)]
    draw.line(pts, fill=color, width=6, joint="curve")
    for x, y in pts:
        draw.ellipse((x - 9, y - 9, x + 9, y + 9), fill=color)


def icon_play(draw, cx, cy, color=WHITE):
    draw.polygon([(cx - 10, cy - 18), (cx - 10, cy + 18), (cx + 22, cy)], fill=color)


def icon_pause(draw, cx, cy, color=WHITE):
    draw.rectangle((cx - 18, cy - 18, cx - 6, cy + 18), fill=color)
    draw.rectangle((cx + 6, cy - 18, cx + 18, cy + 18), fill=color)


def base_screen(title, active="首页"):
    img = Image.new("RGB", (W, H), "#E7ECE9")
    draw = ImageDraw.Draw(img)
    rounded(draw, (70, 44, W - 70, H - 44), 64, fill=PHONE_BG)
    draw.rectangle((120, 44, W - 120, 110), fill="#000000")
    draw.rounded_rectangle((120, 44, W - 120, 120), radius=34, fill="#000000")
    text(draw, (122, 144), "18:26", fill=INK, f=F["sm"])
    draw.rectangle((W - 180, 152, W - 112, 178), outline=INK, width=3)
    draw.rectangle((W - 108, 160, W - 102, 170), fill=INK)
    draw.rectangle((W - 174, 158, W - 124, 172), fill=INK)
    text(draw, (W / 2, 240), title, fill=INK, f=F["h3"], anchor="mm")
    draw.line((92, 294, W - 92, 294), fill="#D8DDD9", width=2)
    return img, draw


def nav(draw, active):
    y = H - 186
    rounded(draw, (90, y, W - 90, H - 64), 36, fill=WHITE, outline="#DCE5E0", width=2)
    items = [("首页", "⌂"), ("拍照", "◎"), ("路线", "⌁"), ("我的", "◉")]
    xs = [180, 360, 540, 720]
    for (label, glyph), x in zip(items, xs):
        fill = GREEN if label == active else MUTED
        text(draw, (x, y + 38), glyph, fill=fill, f=F["h3"], anchor="mm")
        text(draw, (x, y + 82), label, fill=fill, f=F["tab"], anchor="mm")


def pill(draw, x, y, label, fill=GREEN_LIGHT, color=GREEN_DARK):
    tw = draw.textlength(label, font=F["xs"])
    rounded(draw, (x, y, x + tw + 34, y + 44), 22, fill=fill)
    text(draw, (x + 17, y + 9), label, fill=color, f=F["xs"])
    return x + tw + 46


def draw_home():
    img, draw = base_screen("本地历史导览", "首页")
    text(draw, (118, 340), "北京市 · 东城区", f=F["h2"])
    text(draw, (118, 396), "附近历史文化地点与今日路线", fill=MUTED, f=F["body"])
    rounded(draw, (118, 464, 782, 548), 28, fill=WHITE, outline=BORDER, width=2)
    text(draw, (154, 488), "搜索景点、人物、历史事件", fill="#98A2B3", f=F["body"])
    rounded(draw, (118, 590, 782, 864), 30, fill="#DCE8E3", outline="#CADBD4", width=2)
    text(draw, (154, 626), "地图入口保留在首页", f=F["h3"], fill=GREEN_DARK)
    for x, y, label in [(250, 720, "故宫"), (450, 668, "景山"), (585, 770, "什刹海")]:
        draw.ellipse((x - 20, y - 20, x + 20, y + 20), fill=GREEN)
        text(draw, (x, y + 34), label, f=F["xs"], fill=INK, anchor="mm")
    rounded(draw, (610, 800, 750, 840), 20, fill=WHITE)
    text(draw, (632, 808), "查看地图", fill=GREEN, f=F["xs"])
    text(draw, (118, 924), "推荐景点", f=F["h3"])
    places = [("故宫博物院", "1.2km · 皇家文化 · 适合加入路线"), ("天坛公园", "3.8km · 祭祀建筑 · 语音讲解"), ("什刹海", "2.4km · 老北京街区 · 可分享")]
    y = 982
    for name, desc in places:
        rounded(draw, (118, y, 782, y + 126), 24, fill=WHITE, outline=BORDER, width=2)
        rounded(draw, (142, y + 22, 238, y + 104), 18, fill="#D9E8E2")
        text(draw, (266, y + 22), name, f=F["body_b"])
        text(draw, (266, y + 68), desc, fill=MUTED, f=F["sm"])
        y += 146
    nav(draw, "首页")
    return img


def draw_camera():
    img, draw = base_screen("拍照识别", "拍照")
    rounded(draw, (118, 338, 782, 1018), 36, fill="#111827")
    for i in range(4):
        draw.line((118, 450 + i * 120, 782, 450 + i * 120), fill="#27313F", width=2)
    for i in range(4):
        draw.line((250 + i * 120, 338, 250 + i * 120, 1018), fill="#27313F", width=2)
    rounded(draw, (250, 500, 650, 820), 18, fill=None, outline="#F5D46D", width=6)
    text(draw, (450, 474), "对准景点或文物", fill=WHITE, f=F["h3"], anchor="mm")
    fit_text(draw, (170, 860), "AI 会识别可能的文物、景点或建筑细节，并生成历史故事。", 560, fill="#D5DEE8", f=F["body"], max_lines=2)
    rounded(draw, (170, 1058, 730, 1148), 28, fill=WHITE)
    icon_camera(draw, 234, 1104, GREEN)
    text(draw, (288, 1082), "拍摄识别", f=F["body_b"])
    text(draw, (288, 1122), "也可以从相册选择", fill=MUTED, f=F["xs"])
    rounded(draw, (176, 1200, 380, 1268), 28, fill=GREEN_LIGHT)
    text(draw, (278, 1234), "相册", fill=GREEN_DARK, f=F["body_b"], anchor="mm")
    rounded(draw, (520, 1200, 724, 1268), 28, fill=GREEN)
    text(draw, (622, 1234), "识别", fill=WHITE, f=F["body_b"], anchor="mm")
    nav(draw, "拍照")
    return img


def draw_result():
    img, draw = base_screen("识别结果", "拍照")
    rounded(draw, (118, 334, 782, 620), 30, fill="#D6E6DF", outline=BORDER, width=2)
    text(draw, (154, 374), "识别到：故宫角楼", f=F["h2"], fill=GREEN_DARK)
    pill(draw, 154, 444, "置信度 86%")
    pill(draw, 338, 444, "明清建筑")
    fit_text(draw, (154, 518), "角楼位于紫禁城城墙四角，兼具防御、礼制与视觉秩序，是故宫建筑群中辨识度极高的部分。", 560, f=F["sm"], fill=INK, max_lines=3)
    text(draw, (118, 674), "历史故事", f=F["h3"])
    rounded(draw, (118, 730, 782, 1006), 28, fill=WHITE, outline=BORDER, width=2)
    fit_text(draw, (154, 766), "传说角楼设计精巧，九梁十八柱七十二条脊的说法广为流传。准确结构仍应以文物资料为准，但它确实体现了古代木构建筑的复杂营造能力。", 590, f=F["body"], fill=INK)
    rounded(draw, (118, 1044, 782, 1124), 28, fill=GREEN)
    text(draw, (450, 1084), "加入我的路线", fill=WHITE, f=F["body_b"], anchor="mm")
    rounded(draw, (118, 1148, 434, 1228), 28, fill=WHITE, outline="#CFE1DC", width=2)
    text(draw, (276, 1188), "收藏", fill=GREEN, f=F["body_b"], anchor="mm")
    rounded(draw, (466, 1148, 782, 1228), 28, fill=WHITE, outline="#CFE1DC", width=2)
    text(draw, (624, 1188), "分享给朋友", fill=GREEN, f=F["body_b"], anchor="mm")
    nav(draw, "拍照")
    return img


def draw_route():
    img, draw = base_screen("我的路线", "路线")
    text(draw, (118, 344), "北京历史半日游", f=F["h2"])
    text(draw, (118, 398), "可拖拽排序，也可一键分享", fill=MUTED, f=F["body"])
    rounded(draw, (118, 462, 782, 544), 28, fill=GREEN)
    text(draw, (450, 503), "+ 添加地点", fill=WHITE, f=F["body_b"], anchor="mm")
    nodes = [("1", "天安门广场", "建议 45 分钟"), ("2", "故宫博物院", "建议 3 小时"), ("3", "景山公园", "建议 1 小时"), ("4", "什刹海", "建议 1.5 小时")]
    y = 606
    for num, name, sub in nodes:
        rounded(draw, (118, y, 782, y + 132), 26, fill=WHITE, outline=BORDER, width=2)
        draw.ellipse((146, y + 38, 204, y + 96), fill=GREEN_LIGHT)
        text(draw, (175, y + 67), num, fill=GREEN_DARK, f=F["body_b"], anchor="mm")
        text(draw, (230, y + 28), name, f=F["body_b"])
        text(draw, (230, y + 76), sub, fill=MUTED, f=F["sm"])
        text(draw, (726, y + 66), "☰", fill=MUTED, f=F["h3"], anchor="mm")
        y += 154
    rounded(draw, (118, 1238, 424, 1316), 28, fill=WHITE, outline="#CFE1DC", width=2)
    text(draw, (271, 1277), "智能排序", fill=GREEN, f=F["body_b"], anchor="mm")
    rounded(draw, (456, 1238, 782, 1316), 28, fill=GREEN)
    text(draw, (619, 1277), "分享路线", fill=WHITE, f=F["body_b"], anchor="mm")
    nav(draw, "路线")
    return img


def draw_place():
    img, draw = base_screen("地点详情", "首页")
    rounded(draw, (118, 326, 782, 590), 32, fill="#D9E8E2", outline=BORDER, width=2)
    text(draw, (154, 374), "故宫博物院", f=F["h1"], fill=GREEN_DARK)
    pill(draw, 154, 454, "世界文化遗产")
    pill(draw, 376, 454, "皇家文化")
    text(draw, (154, 524), "距离 1.2km · 推荐游览 3-4 小时", fill=MUTED, f=F["body"])
    rounded(draw, (118, 632, 782, 714), 28, fill=GREEN)
    icon_play(draw, 230, 673)
    text(draw, (450, 673), "播放历史讲解", fill=WHITE, f=F["body_b"], anchor="mm")
    text(draw, (118, 774), "核心操作", f=F["h3"])
    actions = [("加入路线", GREEN), ("分享", GREEN_DARK), ("推荐", "#0F8B66"), ("不推荐", RED)]
    x, y = 118, 828
    for label, color in actions:
        rounded(draw, (x, y, x + 312, y + 84), 26, fill=WHITE, outline="#D9E0E7", width=2)
        text(draw, (x + 156, y + 42), label, fill=color, f=F["body_b"], anchor="mm")
        x += 352
        if x > 470:
            x = 118
            y += 112
    rounded(draw, (118, 1092, 782, 1278), 28, fill=WHITE, outline=BORDER, width=2)
    text(draw, (154, 1128), "推荐理由示例", f=F["body_b"])
    fit_text(draw, (154, 1174), "“适合第一次来北京的人，路线清晰，历史信息丰富。”初版建议先审核后公开展示。", 570, f=F["sm"], fill=MUTED)
    nav(draw, "首页")
    return img


def draw_profile():
    img, draw = base_screen("我的", "我的")
    rounded(draw, (118, 340, 782, 556), 32, fill=GREEN)
    draw.ellipse((162, 390, 274, 502), fill=WHITE)
    icon_person(draw, 218, 452, GREEN)
    text(draw, (310, 390), "编辑个人资料", f=F["h2"], fill=WHITE)
    text(draw, (310, 452), "昵称、头像、出行偏好", f=F["body"], fill="#E9FBF7")
    rounded(draw, (588, 456, 742, 514), 22, fill=WHITE)
    text(draw, (665, 485), "进入", fill=GREEN, f=F["sm"], anchor="mm")
    items = [("我的路线", "管理自定义路线并分享"), ("模式设置", "普通 / 无障碍出行 / 语音大字"), ("隐私与授权", "位置、相册、头像权限管理"), ("帮助反馈", "提交识别错误或体验建议")]
    y = 622
    for title, sub in items:
        rounded(draw, (118, y, 782, y + 122), 24, fill=WHITE, outline=BORDER, width=2)
        text(draw, (154, y + 26), title, f=F["body_b"])
        text(draw, (154, y + 72), sub, fill=MUTED, f=F["sm"])
        text(draw, (738, y + 61), "›", fill=MUTED, f=F["h2"], anchor="mm")
        y += 144
    nav(draw, "我的")
    return img


def draw_modes():
    img, draw = base_screen("选择模式", "我的")
    text(draw, (118, 346), "选择适合你的导览方式", f=F["h2"])
    text(draw, (118, 402), "模式入口放到“我的”，首页保持轻量。", fill=MUTED, f=F["body"])
    cards = [
        ("普通模式", "内容不变，适合多数游客", "person", GREEN),
        ("无障碍出行", "少台阶、短步行、休息点优先", "route", GOLD),
        ("语音大字模式", "保留视障能力，并合并无障碍路线", "audio", "#111111"),
    ]
    y = 506
    for title, sub, icon, color in cards:
        rounded(draw, (118, y, 782, y + 210), 32, fill=WHITE, outline=BORDER, width=2)
        draw.ellipse((154, y + 44, 270, y + 160), fill=GREEN_LIGHT if color != "#111111" else "#EFEFEF")
        if icon == "person":
            icon_person(draw, 212, y + 100, color)
        elif icon == "route":
            icon_route(draw, 212, y + 104, color)
        else:
            draw.ellipse((178, y + 72, 246, y + 140), fill=color)
            icon_play(draw, 212, y + 106, WHITE)
        text(draw, (304, y + 54), title, f=F["h3"])
        fit_text(draw, (304, y + 104), sub, 430, fill=MUTED, f=F["sm"], max_lines=2)
        y += 244
    nav(draw, "我的")
    return img


def draw_architecture():
    img = Image.new("RGB", (1400, 900), "#F6F8FA")
    draw = ImageDraw.Draw(img)
    text(draw, (70, 62), "改版后的四条主线", f=font(48, "bold"), fill=INK)
    text(draw, (70, 122), "首页负责发现，拍照负责识别，路线负责计划，我的负责资料与模式。", f=font(28), fill=MUTED)
    centers = [(220, 330), (520, 330), (820, 330), (1120, 330)]
    titles = ["首页发现", "拍照识别", "我的路线", "我的"]
    descs = ["推荐景点\n地图入口\n全部/附近", "景点文物识别\n历史故事\n加入路线", "拖拽排序\n智能排序\n分享朋友", "头像昵称\n模式设置\n隐私授权"]
    colors = [GREEN, "#0F8B66", BLUE, GOLD]
    for (cx, cy), title_s, desc, c in zip(centers, titles, descs, colors):
        rounded(draw, (cx - 120, cy - 92, cx + 120, cy + 92), 28, fill=WHITE, outline="#D0D7DE", width=2)
        draw.ellipse((cx - 42, cy - 64, cx + 42, cy + 20), fill=c)
        text(draw, (cx, cy + 56), title_s, f=font(28, "bold"), fill=INK, anchor="mm")
        if cx < 1120:
            draw.line((cx + 130, cy, cx + 170, cy), fill="#8C959F", width=5)
            draw.polygon([(cx + 170, cy), (cx + 154, cy - 10), (cx + 154, cy + 10)], fill="#8C959F")
    y = 560
    for i, (title_s, desc, c) in enumerate(zip(titles, descs, colors), start=1):
        x = 88 + (i - 1) * 320
        rounded(draw, (x, y, x + 278, y + 238), 24, fill=WHITE, outline="#D0D7DE", width=2)
        text(draw, (x + 24, y + 26), f"{i}. {title_s}", f=font(26, "bold"), fill=c)
        yy = y + 78
        for line in desc.split("\n"):
            text(draw, (x + 24, yy), line, f=font(24), fill=INK)
            yy += 42
    return img


def save_images():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    specs = [
        ("01_home_navigation.png", draw_home(), "首页信息架构：底部导航改为首页、拍照、路线、我的，地图入口留在首页。"),
        ("02_camera_identify.png", draw_camera(), "拍照识别入口：支持拍摄与相册选择，突出文物/景点识别。"),
        ("03_ai_result.png", draw_result(), "AI 识别结果：输出历史故事，并支持加入路线、收藏和分享。"),
        ("04_custom_route.png", draw_route(), "我的路线：支持添加地点、拖拽排序、智能排序和分享。"),
        ("05_place_detail_actions.png", draw_place(), "地点详情：加入路线、分享、推荐/不推荐和单按钮语音。"),
        ("06_profile_edit.png", draw_profile(), "我的页面：编辑头像昵称，保留路线、模式、隐私等关键入口。"),
        ("07_mode_selection.png", draw_modes(), "模式选择：普通、无障碍出行、语音大字三模式，普通模式使用人物图标。"),
        ("08_architecture.png", draw_architecture(), "整体信息架构：四条主线把功能收束清楚。"),
    ]
    paths = []
    for name, img, caption in specs:
        path = OUT_DIR / name
        img.save(path, quality=95)
        paths.append((path, caption))
    return paths


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, content="", style=None, size=11, bold=False, color="000000", align=None):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    r = p.add_run(content)
    set_font(r, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, content):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(content)
    set_font(r, size=11)


def build_doc(paths):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "STHeiti")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("本地历史导览小程序改版方案与界面图")
    set_font(run, size=22, bold=True, color="0B2545")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("围绕拍照识别、自定义路线、分享、个人资料与无障碍模式的产品升级建议")
    set_font(r, size=11, color="555555")

    doc.add_paragraph()
    add_paragraph(doc, "一、改版目标", style="Heading 1")
    for item in [
        "把新增功能收束为四条清晰主线：发现地点、拍照识别、我的路线、个人与无障碍设置。",
        "减少首页冗余入口，避免“切换模式”“地图”“收藏”等功能互相抢位置。",
        "把拍照识别做成新的核心能力，让用户可以对景点、建筑细节、文物进行 AI 识别和故事化讲解。",
        "把路线功能从推荐路线升级为用户可创建、可排序、可分享的旅行计划。",
    ]:
        add_bullet(doc, item)

    add_paragraph(doc, "二、建议的信息架构", style="Heading 1")
    add_paragraph(doc, "底部导航建议调整为：首页 / 拍照 / 路线 / 我的。地图不再作为独立底栏项，而是保留在首页与地点详情中，作为空间定位和导航辅助入口。")
    doc.add_picture(str(paths[-1][0]), width=Inches(6.3))
    cap = doc.add_paragraph(paths[-1][1])
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(cap.runs[0], size=9, color="555555")

    add_paragraph(doc, "三、功能取舍与风险", style="Heading 1")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2100, 3900, 3360]
    headers = ["模块", "建议做法", "原因"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(header)
        set_font(r, size=10, bold=True, color="1F4D78")
        set_cell_shading(cell, "F2F4F7")
        set_cell_width(cell, widths[i])
    rows = [
        ("拍照识别", "作为独立底栏入口", "新功能高频且有差异化，不应藏在地图或详情里。"),
        ("地图", "保留在首页和详情中", "首页推荐和地图并不重复，一个负责发现，一个负责空间判断。"),
        ("评价理由", "先审核后公开，或第一版只展示统计", "公开文字内容需要内容安全、举报和反垃圾机制。"),
        ("注册资料", "做成编辑资料，不做强注册", "微信小程序用户更适合轻量授权和渐进完善资料。"),
        ("语音按钮", "合并为一个播放/暂停/继续按钮", "降低操作复杂度，尤其适合视障和行动不便用户。"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(value)
            set_font(r, size=9.5)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cells[i], widths[i])

    add_paragraph(doc, "四、界面方案图", style="Heading 1")
    for path, caption in paths[:-1]:
        add_paragraph(doc, caption.split("：")[0], style="Heading 2")
        add_paragraph(doc, caption)
        doc.add_picture(str(path), width=Inches(3.15))
        cap = doc.add_paragraph(path.name)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cap.runs[0], size=8.5, color="666666")

    add_paragraph(doc, "五、实施优先级", style="Heading 1")
    add_paragraph(doc, "第一阶段建议优先落地：拍照识别、地点加入路线、路线排序与分享、语音单按钮、个人资料编辑、首页移除切换模式。")
    add_paragraph(doc, "第二阶段再扩展：推荐/不推荐理由公开、收藏分享、更多无障碍偏好、全部景点与附近景点的组合筛选。")

    doc.save(DOCX_PATH)


def main():
    paths = save_images()
    build_doc(paths)
    print(DOCX_PATH)
    for path, _ in paths:
        print(path)


if __name__ == "__main__":
    main()
